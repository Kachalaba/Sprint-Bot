"""Export utilities for personal best and Sum of Best analytics."""

from __future__ import annotations

import asyncio
import csv
import io
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook

from services.pb_service import get_sob, get_total_pb_attempt
from utils.logger import get_logger

logger = get_logger(__name__)

_DEFAULT_DB_PATH = Path("data/results.db")


@dataclass(frozen=True, slots=True)
class ExportRow:
    """Single row describing PB/SoB analytics."""

    athlete_id: int
    athlete_name: str
    stroke: str
    distance: int
    segment_index: int
    pb_split: float | None
    sob_split: float | None
    pb_total: float | None
    sob_total: float | None


class ExportService:
    """Prepare exports with advanced analytics."""

    def __init__(self, db_path: Path | str = _DEFAULT_DB_PATH) -> None:
        self._path = Path(db_path)

    async def export_pb_sob(
        self,
        athlete_ids: Sequence[int],
        *,
        stroke: str | None = None,
        distance: int | None = None,
        fmt: str = "csv",
    ) -> bytes:
        """Export PB and SoB data for provided athletes."""

        rows = await asyncio.to_thread(
            self._collect_rows, tuple(athlete_ids), stroke, distance
        )
        if fmt.lower() == "csv":
            return self._to_csv(rows)
        if fmt.lower() in {"xlsx", "excel"}:
            return self._to_excel(rows)
        raise ValueError("Unsupported export format")

    def _collect_rows(
        self,
        athlete_ids: Sequence[int],
        stroke: str | None,
        distance: int | None,
    ) -> list[ExportRow]:
        if not athlete_ids:
            return []
        filtered_stroke = stroke.strip().lower() if stroke else None
        with sqlite3.connect(self._path) as conn:
            conn.row_factory = sqlite3.Row
            placeholders = ",".join("?" for _ in athlete_ids)
            query = (
                "SELECT DISTINCT athlete_id, athlete_name, stroke, distance "
                "FROM results WHERE athlete_id IN ({ids})".format(ids=placeholders)
            )
            args: list[object] = [int(aid) for aid in athlete_ids]
            if filtered_stroke is not None:
                query += " AND stroke = ?"
                args.append(filtered_stroke)
            if distance is not None:
                query += " AND distance = ?"
                args.append(int(distance))
            query += " ORDER BY athlete_id, stroke, distance"
            pairs = conn.execute(query, args).fetchall()

        rows: list[ExportRow] = []
        for row in pairs:
            aid = int(row["athlete_id"])
            name = str(row["athlete_name"] or f"ID {aid}")
            stroke_value = str(row["stroke"])
            distance_value = int(row["distance"])
            pb = get_total_pb_attempt(
                aid, stroke_value, distance_value, db_path=self._path
            )
            sob = get_sob(aid, stroke_value, distance_value, db_path=self._path)
            segments = pb.segments if pb else tuple()
            sob_segments = sob.segments
            max_len = max(len(segments), len(sob_segments))
            for idx in range(max_len or 1):
                pb_split = segments[idx] if idx < len(segments) else None
                sob_split = sob_segments[idx] if idx < len(sob_segments) else None
                rows.append(
                    ExportRow(
                        athlete_id=aid,
                        athlete_name=name,
                        stroke=stroke_value,
                        distance=distance_value,
                        segment_index=idx + 1,
                        pb_split=pb_split,
                        sob_split=sob_split,
                        pb_total=pb.total if pb else None,
                        sob_total=sob.total,
                    )
                )
        return rows

    def _to_csv(self, rows: Iterable[ExportRow]) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(
            [
                "athlete_id",
                "athlete_name",
                "stroke",
                "distance",
                "segment_index",
                "pb_split",
                "sob_split",
                "pb_total",
                "sob_total",
            ]
        )
        for row in rows:
            writer.writerow(
                [
                    row.athlete_id,
                    row.athlete_name,
                    row.stroke,
                    row.distance,
                    row.segment_index,
                    row.pb_split if row.pb_split is not None else "",
                    row.sob_split if row.sob_split is not None else "",
                    row.pb_total if row.pb_total is not None else "",
                    row.sob_total if row.sob_total is not None else "",
                ]
            )
        return buffer.getvalue().encode("utf-8")

    def _to_excel(self, rows: Iterable[ExportRow]) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "PB_SoB"
        headers = [
            "athlete_id",
            "athlete_name",
            "stroke",
            "distance",
            "segment_index",
            "pb_split",
            "sob_split",
            "pb_total",
            "sob_total",
        ]
        sheet.append(headers)
        for row in rows:
            sheet.append(
                [
                    row.athlete_id,
                    row.athlete_name,
                    row.stroke,
                    row.distance,
                    row.segment_index,
                    row.pb_split,
                    row.sob_split,
                    row.pb_total,
                    row.sob_total,
                ]
            )
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


__all__ = ["ExportRow", "ExportService"]
